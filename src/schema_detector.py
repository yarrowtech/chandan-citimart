"""Workbook schema detection and alias-based field mapping."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from config.column_aliases import COLUMN_ALIASES
from config.settings import STORE_SHEET_PREFIXES


def normalize_name(value: object) -> str:
    """Normalize case, whitespace, and punctuation for tolerant matching."""
    text = str(value or "").strip().casefold()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def detect_columns(columns: Iterable[object]) -> dict[str, str]:
    """Return canonical-to-source mappings, preferring exact normalized aliases."""
    normalized = {normalize_name(column): str(column) for column in columns}
    mapping: dict[str, str] = {}
    for canonical, aliases in COLUMN_ALIASES.items():
        for alias in [canonical, *aliases]:
            match = normalized.get(normalize_name(alias))
            if match is not None:
                mapping[canonical] = match
                break
    return mapping


def store_code_for_sheet(sheet_name: str) -> str | None:
    """Resolve a sheet name to a configured store code."""
    normalized = normalize_name(sheet_name).replace(" ", "")
    for code, prefixes in STORE_SHEET_PREFIXES.items():
        if any(normalized.startswith(prefix.casefold()) for prefix in prefixes):
            return code
    return None


def sheet_role(sheet_name: str) -> str | None:
    """Infer the current workbook's pair role from its numeric suffix."""
    match = re.search(r"(\d+)\s*$", sheet_name)
    if not match:
        return None
    return {"1": "detail", "2": "daily"}.get(match.group(1))


@dataclass(frozen=True)
class SheetInspection:
    name: str
    rows: int
    columns: int
    headers: tuple[str, ...]
    store_code: str | None
    role: str | None
    mapping: dict[str, str]


def inspect_workbook(path: Path) -> list[SheetInspection]:
    """Inspect dimensions and first-row headers without altering the workbook."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    inspections: list[SheetInspection] = []
    try:
        for worksheet in workbook.worksheets:
            headers = tuple(
                str(cell.value).strip() if cell.value is not None else ""
                for cell in next(worksheet.iter_rows(min_row=1, max_row=1))
            )
            inspections.append(
                SheetInspection(
                    name=worksheet.title,
                    rows=max(worksheet.max_row - 1, 0),
                    columns=worksheet.max_column,
                    headers=headers,
                    store_code=store_code_for_sheet(worksheet.title),
                    role=sheet_role(worksheet.title),
                    mapping=detect_columns(headers),
                )
            )
    finally:
        workbook.close()
    return inspections

