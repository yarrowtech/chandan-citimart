"""Load governed KPI colour bands from the workbook reference sheet."""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from config.kpi_thresholds import KPIColorRule
from src.schema_detector import normalize_name


COLOR_SHEET_NAMES = {
    "colour format",
    "colour formatting",
    "color format",
    "color formatting",
}

KPI_NAME_MAP = {
    "atv": "atv",
    "rpv": "rpv",
    "basket size": "basket_size",
    "busket size": "basket_size",
    "conversion": "conversion",
    "conversion percentage": "conversion",
    "achievement": "achievement",
    "achievement percentage": "achievement",
    "achievemnet percentage": "achievement",
}

PERCENTAGE_KPIS = {"conversion", "achievement"}


def is_color_rule_sheet(sheet_name: str) -> bool:
    return normalize_name(sheet_name) in COLOR_SHEET_NAMES


def _numbers(value: object) -> list[float]:
    return [
        float(item)
        for item in re.findall(r"-?\d+(?:\.\d+)?", str(value or ""))
    ]


def load_color_rules(
    path: str | Path,
) -> tuple[dict[str, KPIColorRule], str | None, list[str]]:
    """Return validated colour rules, controlling sheet name, and warnings."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    warnings: list[str] = []
    try:
        worksheet = next(
            (
                sheet
                for sheet in workbook.worksheets
                if is_color_rule_sheet(sheet.title)
            ),
            None,
        )
        if worksheet is None:
            return {}, None, [
                "KPI colour coding is neutral because COLOUR FORMATTING was not found."
            ]

        rows = worksheet.iter_rows(values_only=True)
        headers = [normalize_name(value) for value in next(rows, ())]
        required = {"kpi name", "red", "yellow", "green"}
        if not required.issubset(headers):
            return {}, worksheet.title, [
                f"{worksheet.title}: expected KPI NAME, RED, YELLOW, and GREEN columns."
            ]
        positions = {name: headers.index(name) for name in required}
        rules: dict[str, KPIColorRule] = {}
        for row_number, row in enumerate(rows, start=2):
            label = row[positions["kpi name"]]
            normalized_label = normalize_name(label)
            if not normalized_label:
                continue
            metric = KPI_NAME_MAP.get(normalized_label)
            if metric is None:
                warnings.append(
                    f"{worksheet.title} row {row_number}: unsupported KPI {label!r}."
                )
                continue
            red_text = str(row[positions["red"]] or "").strip()
            yellow_text = str(row[positions["yellow"]] or "").strip()
            green_text = str(row[positions["green"]] or "").strip()
            red_values = _numbers(red_text)
            yellow_values = _numbers(yellow_text)
            green_values = _numbers(green_text)
            if not red_values or len(yellow_values) < 2 or not green_values:
                warnings.append(
                    f"{worksheet.title} row {row_number}: incomplete colour ranges."
                )
                continue
            lower, upper = yellow_values[0], yellow_values[-1]
            if (
                lower >= upper
                or red_values[-1] != lower
                or green_values[0] != upper
            ):
                warnings.append(
                    f"{worksheet.title} row {row_number}: colour boundaries do not reconcile."
                )
                continue
            if metric in PERCENTAGE_KPIS:
                lower /= 100.0
                upper /= 100.0
            rules[metric] = KPIColorRule(
                metric=metric,
                label=str(label),
                lower_bound=lower,
                upper_bound=upper,
                red_text=red_text,
                yellow_text=yellow_text,
                green_text=green_text,
                source=worksheet.title,
            )
        return rules, worksheet.title, warnings
    finally:
        workbook.close()
